#!/usr/bin/env python3

'''Main script for PSAN evaluation. Please read the main() at the bottom for more details.
Note: "analyzer" refers to the PSAN tool
'''

import shutil
import subprocess
import os
import pathlib
import typing
import argparse
import dataclasses
import datetime
import json
import collections
import re
import threading
import concurrent.futures
import openpyxl
import yaml

def env(name : str) -> str:
  '''Get the value of an environment variable, or raise an exception if it is not set.'''
  val = os.environ.get(name)
  if val is None:
    raise RuntimeError("Environment variable " + name + " is not set")
  return val

src_ir_path = env("EVAL_IRDIR") # whole-program IR files should be placed in {src_ir_path}/${program_set}/${program_name}
root_output_path = env("EVAL_EXEDIR") # where to place final executables and necessary files

run_script_base_path = root_output_path # where the run script will be placed
run_data_rootdir = root_output_path + '/data'
run_log_rootdir = root_output_path + '/log'
run_perfstat_log_rootdir = root_output_path + '/perfstat-log'

run_script_name_integrated = 'run_integrated.sh'  # program output integrated in terminal (for FPGA run)
run_script_name_split =      'run_split.sh'       # program output in sparate files (for QEMU run)
run_script_name_regression = 'run_regression.sh'  # for QEMU based regression testing (only execute once, split output)
run_script_name_perfstat =   'run_perfstat.sh'    # for performance diagnostics using perf stat
run_script_name_perfstat_1 = 'run_perfstat_1.sh'  # same, but just once per program
run_script_name_rt_stats =   'run_rt_stats.sh'    # for runtime statistics collection
run_script_name_split_fixms = 'run_split_fixms.sh' # same as run_split.sh, but only run if output file is missing
makefile_regression_outpath = 'Makefile_regression' # replacing run_script_name_regression for parallel testing

static_stats_output_path                = 'psan_static_stats.csv'
static_stats_output_path_xlsx           = 'psan_static_stats.xlsx'
runtime_time_stats_output_path          = 'psan_runtime_stats_time.csv'
runtime_time_stats_digested_output_path = 'psan_runtime_stats_time_digested.csv'
runtime_psan_stats_output_path          = 'psan_runtime_stats_psan.csv'
runtime_perf_stats_output_path          = 'psan_perf_stats.csv'

ir_suffix_in = 'bc' # read in bitcode files
ir_suffix_temp = 'll' # write text ir file out

run_script_integrated_log_start = '>>>>>>>>>>>>>>>> PSAN-EVAL START: '
run_script_integrated_log_end =   '>>>>>>>>>>>>>>>> PSAN-EVAL END: '

# [benchmark set name] -> [list of programs]
enabled_programs_dict = {
  "specint2006" : [
    #'400.perlbench',
    '401.bzip2',
    #'403.gcc',
    '429.mcf',
    '445.gobmk',
    '456.hmmer',
    '458.sjeng',
    '462.libquantum',
    '464.h264ref',

    # C++ programs not supported
    #'471.omnetpp',
    #'473.astar',
    #'483.xalancbmk',
  ],
  "spec2017" : [
    #"500.perlbench_r",
    #"502.gcc_r",
    "505.mcf_r",
    "519.lbm_r",
    "525.x264_r",
    "538.imagick_r",
    "544.nab_r",
    "557.xz_r",
  ],
  "olden" : [
    'bh',
    'bisort',
    'em3d',
    'health',
    'mst',
    'perimeter',
    'power',
    'treeadd',
    'tsp',
    'voronoi',
  ],
  "ptrdist" : [
    'anagram',
    'bc',
    'ft',
    'ks',
    'yacr2',
  ],
  "misc" : [
  #  'coremark',
  #  'coremark-mifp',
    'nginx',
    'pureftpd',
  ],
}
disabled_program_set_names : list[str] = [
  "spec2017", # disable spec2017 by default, since we cannot distribute them
  "specint2006",
  #"olden",
  "ptrdist",
  "misc",
]
for disabled_set in disabled_program_set_names:
  del enabled_programs_dict[disabled_set]

cpp_programs_list = [
  '471.omnetpp',
  '473.astar',
  '483.xalancbmk',
]

# each node: (command, overrides)
compilation_command_info = (
  ["-lm", "-O3", "{in}", "-o", "{out}"], # defaults
  {
    "misc": (
      ["-lm", "-O3", "{in}", "-o", "{out}"], # defaults
      {
        "nginx": (["-lcrypt", "-lpcre", "-lz", "-ldl", "-lpthread", "-O3", "{in}", "-o", "{out}"], None),
        "pureftpd": (["-lcrypt", "-O3", "{in}", "-o", "{out}"], None),
      }
    )
  }) # overrides


def parse_compilation_command_info(path):
  global compilation_command_info

  def convert(node):
    if isinstance(node, dict):
      default = node.get("default")
      override = node.get("override")
      if override is not None:
        # Recursively convert each override entry
        override_converted = {k: convert(v) for k, v in override.items()}
      else:
        override_converted = None
      return (default, override_converted)
    else:
      # Leaf node
      return node, None

  with open(path, "r", encoding="utf-8") as f:
    data = yaml.safe_load(f)
    compilation_command_info = convert(data)

PSANEVAL_OVERRIDE_COMPILE_COMMAND_PATH = os.environ.get("PSANEVAL_OVERRIDE_COMPILE_COMMAND", "psaneval_override_compile_command.yml")
if os.path.isfile(PSANEVAL_OVERRIDE_COMPILE_COMMAND_PATH):
  parse_compilation_command_info(PSANEVAL_OVERRIDE_COMPILE_COMMAND_PATH)

compile_only_programs = [
  "misc/nginx",
  "misc/pureftpd",
]

program_command_dict = {
  "olden/bh":         "16384 8",    #"4096 1",
  "olden/bisort":     "2000000 1",  #"250000 1",
  "olden/em3d":       "10000 500 75 1", #"2000 100 75 1",
  "olden/health":     "6 400 4",    #"5 500 4",
  "olden/mst":        "4096 0",     #"1024 0",
  "olden/perimeter":  "12 0",       #"9 0",
  "olden/power":      "",
  "olden/treeadd":    "25 1 1",     #"20 1 1",
  "olden/tsp":        "2000000 1",  #"100000 1",
  "olden/voronoi":    "60000 1",    #"20000 1",

  "ptrdist/anagram":  "{file_path}/words < {file_path}/input.in",
  "ptrdist/bc":       "< {file_path}/primes.b",
  "ptrdist/ft":       "5000 6000",
  "ptrdist/ks":       "{file_path}/KL-1.in",
  "ptrdist/yacr2":    "{file_path}/input2.in",

  # 400.perlbench is still buggy and can crash in some inputs
  #"specint2006/400.perlbench": "-I{file_path}/lib {file_path}/checkspam.pl 2500 5 25 11 150 1 1 1 1",
  #"specint2006/400.perlbench": "-I{file_path}/lib {file_path}/diffmail.pl 4 800 10 17 19 300",
  #"specint2006/400.perlbench": "-I{file_path}/lib {file_path}/splitmail.pl 1600 12 26 16 4500",

  #"specint2006/401.bzip2": "{file_path}/input.source 280 > {file_path}/input.source.out",
  #"specint2006/401.bzip2": "{file_path}/chicken.jpg 30 > {file_path}/chicken.jpg.out",
  #"specint2006/401.bzip2": "{file_path}/liberty.jpg 30 > {file_path}/liberty.jpg.out",
  #"specint2006/401.bzip2": "{file_path}/input.program 280 > {file_path}/input.program.out",
  "specint2006/401.bzip2": "{file_path}/text.html 280 > /dev/null",
  #"specint2006/401.bzip2": "{file_path}/input.combined 200 > {file_path}/input.combined.out",

  # gcc does not work yet because of unknown reasons
  #"specint2006/403.gcc":  "{file_path}/166.i -o {file_path}/166.s",
  #"specint2006/403.gcc":  "{file_path}/200.i -o {file_path}/200.s",
  #"specint2006/403.gcc":  "{file_path}/c-typeck.i -o {file_path}/c-typeck.s",
  #"specint2006/403.gcc":  "{file_path}/cp-decl.i -o {file_path}/cp-decl.s",
  #"specint2006/403.gcc":  "{file_path}/expr.i -o {file_path}/expr.s",
  #"specint2006/403.gcc":  "{file_path}/expr2.i -o {file_path}/expr2.s",
  #"specint2006/403.gcc":  "{file_path}/g23.i -o {file_path}/g23.s",
  #"specint2006/403.gcc":  "{file_path}/s04.i -o {file_path}/s04.s",
  #"specint2006/403.gcc":  "{file_path}/scilab.i -o {file_path}/scilab.s",

  "specint2006/429.mcf": "{file_path}/inp.in",

  "specint2006/445.gobmk": "--quiet --mode gtp < {file_path}/13x13.tst",
  #"specint2006/445.gobmk": "--quiet --mode gtp < {file_path}/nngs.tst",
  #"specint2006/445.gobmk": "--quiet --mode gtp < {file_path}/score2.tst",
  #"specint2006/445.gobmk": "--quiet --mode gtp < {file_path}/trevorc.tst",
  #"specint2006/445.gobmk": "--quiet --mode gtp < {file_path}/trevord.tst",

  #"specint2006/456.hmmer": "{file_path}/nph3.hmm {file_path}/swiss41",
  "specint2006/456.hmmer": "--fixed 0 --mean 500 --num 500000 --sd 350 --seed 0 {file_path}/retro.hmm",

  "specint2006/458.sjeng": "{file_path}/ref.txt",

  "specint2006/462.libquantum": "1397 8",

  # 464.h264ref is still buggy and can crash in some inputs
  "specint2006/464.h264ref": "-d {file_path}/foreman_ref_encoder_baseline.cfg",
  #"specint2006/464.h264ref": "-d {file_path}/foreman_ref_encoder_main.cfg",
  #"specint2006/464.h264ref": "-d {file_path}/sss_encoder_main.cfg",

  "spec2017/500.perlbench_r": "-I{file_path}/lib {file_path}/checkspam.pl 2500 5 25 11 150 1 1 1 1",
  #"spec2017/500.perlbench_r": "-I{file_path}/lib {file_path}/diffmail.pl 4 800 10 17 19 300",
  #"spec2017/500.perlbench_r": "-I{file_path}/lib {file_path}/splitmail.pl 6400 12 26 16 100 0",

  "spec2017/505.mcf_r": "{file_path}/inp.in",

  "spec2017/525.x264_r": "--pass 1 --stats {file_path}/x264_stats.log --bitrate 1000 --frames 1000 -o {file_path}/BuckBunny_New.264 {file_path}/BuckBunny.yuv 1280x720",
  # --pass 2 --stats x264_stats.log --bitrate 1000 --dumpyuv 200 --frames 1000 -o BuckBunny_New.264 BuckBunny.yuv 1280x720
  # --seek 500 --dumpyuv 200 --frames 1250 -o BuckBunny_New.264 BuckBunny.yuv 1280x720

  "spec2017/557.xz_r": "{file_path}/cld.tar.xz 160 19cf30ae51eddcbefda78dd06014b4b96281456e078ca7c13e1c0c9e6aaea8dff3efb4ad6b0456697718cede6bd5454852652806a657bb56e07d61128434b474 59796407 61004416 6",
  #"spec2017/557.xz_r": "{file_path}/cpu2006docs.tar.xz 250 055ce243071129412e9dd0b3b69a21654033a9b723d874b2015c774fac1553d9713be561ca86f74e4f16f22e664fc17a79f30caa5ad2c04fbc447549c2810fae 23047774 23513385 6e",
  #"spec2017/557.xz_r": "{file_path}/input.combined.xz 250 a841f68f38572a49d86226b7ff5baeb31bd19dc637a922a972b2e6d1257a890f6a544ecab967c313e370478c74f760eb229d4eef8a8d2836d233d3e9dd1430bf 40401484 41217675 7",

  "spec2017/519.lbm_r": "3000 {file_path}/reference.dat 0 0 {file_path}/100_100_130_ldc.of", # reference.dat is output

  "spec2017/538.imagick_r": r"-limit disk 0 {file_path}/refrate_input.tga -edge 41 -resample 181% -emboss 31 -colorspace YUV -mean-shift 19x19+15% -resize 30% {file_path}/refrate_output.tga",

  "spec2017/544.nab_r": "1am0 1122214447 122", # 1am0 must be relative path

  #"misc/coremark":        "0 0 0 1000",
  #"misc/coremark-mifp":   "0 0 0 1000",
}

toolpath = env("TOOLPATH") # directory where the repo is placed
# program set -> program name -> file path override
# by default, the filepaths are under <run_data_rootdir>/<program_set_name>/<program_name>
# if we want to change it for some programs, we can specify it here
program_filepath_overrides : dict[str, dict[str, str] | str] = {
  "ptrdist": os.path.join(toolpath, "benchsrc/ptrdist/ptrdist-1.1"),
  "specint2006": os.path.join(toolpath, "specdata"),
  "spec2017": os.path.join(toolpath, "spec2017data"),
}

timer_command_prefix_list = ["/usr/bin/time", "-v"] #"/usr/bin/time -al"
timer_command_prefix = ' '.join(timer_command_prefix_list)
perfstat_command_prefix = "perf stat -e cache-references,cache-misses,cycles,instructions,branches,faults,migrations -dd --no-big-num"

# [suffix] -> <Analyzer command: {psan} {in} {out}>
psan_common_flags_base = ["--psan-no-dump", "--psan-no-debug", "--psan-verify-output"]
psan_common_flags = psan_common_flags_base + ["--psan-enable-conservative-oob"]

@dataclasses.dataclass
class ExternalToolPipeline:
  instrument_cmd_env : str # environment variable to find executable/command that does IR instrumentation
  linking_flags_env : str # environment variable for additional linking flags
  require_optimized_input : bool = False # whether the input IR should be optimized before instrumentation

  @staticmethod
  def _unpack_env_var(env_name: str) -> list[str] | None:
    value = os.environ.get(env_name)
    if value is None:
      return None
    # Split the value by spaces and return as a list
    return value.split()

  def get_cmd(self) -> list[str] | None:
    return self._unpack_env_var(self.instrument_cmd_env)
  def get_ldflags(self) -> list[str] | None:
    return self._unpack_env_var(self.linking_flags_env)

  def validate(self) -> bool:
    cmd = self.get_cmd()
    if cmd is None:
      print(self.instrument_cmd_env + " is not set")
      return False
    # the first element should be an executable
    # check if the file exists
    if len(cmd) == 0:
      # may not need this step (e.g., sanitizer)
      return True
    if not os.path.isfile(cmd[0]):
      print(self.instrument_cmd_env + ": Executable " + cmd[0] + " does not exist")
      return False
    return True

  def get_complete_commands(self, ir_file : str, instrumented_ir_out : str):
    if cmds := self.get_cmd():
      return cmds + [ir_file, "-o", instrumented_ir_out]
    return get_disaasemble_command(ir_file, instrumented_ir_out)

enabled_modes = {
  "baseline" : [], # empty for plain copy
  #"spatial":        psan_common_flags + ["--psan-mss=true",  "--psan-mts=false", "{in}", "-o", "{out}"],
  #"spatial-oob":    psan_common_flags + ["--psan-mss=true",  "--psan-mts=false", "{in}", "-o", "{out}", "--psan-disable-eptg"],
  #"temporal":       psan_common_flags + ["--psan-mss=false", "--psan-mts=true", "{in}", "-o", "{out}"],
  #"temporal-oob":   psan_common_flags + ["--psan-mss=false", "--psan-mts=true", "{in}", "-o", "{out}", "--psan-disable-eptg"],
  "full":           psan_common_flags + ["--psan-mss=true",  "--psan-mts=true", "{in}", "-o", "{out}"],
  "full-oob":       psan_common_flags + ["--psan-mss=true",  "--psan-mts=true", "{in}", "-o", "{out}", "--psan-disable-eptg"],
  #"full-nchk":      psan_common_flags + ["--psan-mss=true",  "--psan-mts=true", "--psan-noop-check", "{in}", "-o", "{out}"],
  #"full-oob-nchk":  psan_common_flags + ["--psan-mss=true",  "--psan-mts=true", "--psan-noop-check", "{in}", "-o", "{out}", "--psan-disable-eptg"],
  #"full-stat":      psan_common_flags + ["--psan-mss=true",  "--psan-mts=true", "--psan-collect-runtime-stat", "{in}", "-o", "{out}"],
  #"full-aggr":      psan_common_flags_base + ["--psan-mss=true", "--psan-mts=true", "{in}", "-o", "{out}"],
  #"full-oob-aggr":  psan_common_flags_base + ["--psan-mss=true", "--psan-mts=true", "{in}", "-o", "{out}", "--psan-disable-eptg"],
  #"asan": ExternalToolPipeline(instrument_cmd_env="ASAN_EXEC", linking_flags_env="ASAN_LDFLAGS", require_optimized_input=True),
  "softboundcets": ExternalToolPipeline(instrument_cmd_env="SOFTBOUNDCETS_EXEC", linking_flags_env="SOFTBOUNDCETS_LDFLAGS", require_optimized_input=True),
}
static_stat_mode = "full" # we use the output in this mode for the compile-time stats
dynamic_time_stat_modes = [ # we use the output of time command on these mode(s); 'baseline' MUST be the first entry
  "baseline",
  "spatial",
  "spatial-oob",
  "temporal",
  "temporal-oob",
  "full",
  "full-oob",
  "full-nchk",
  "full-oob-nchk",
  "full-aggr",
  "full-oob-aggr",
  "asan",
  "softboundcets",
]
dynamic_custom_stat_modes = ["full-stat"] #['stats'] # we extract the runtime stats on these mode(s)

psan_path = env("PSAN_EXEC") # path to the PSAN executable
psan_wpa = env("PSAN_WPA") # path to the WPA executable
llvmdis_path = env("LLVM_DIS") # path to the llvm-dis executable
clang_path = env("CLANG") # path to the clang executable
clangpp_path = clang_path + "++"
temp_ir_path = env("EVAL_TMPIR_DIR") # where we place the temporary IR files
num_repeat_execution = 5 # how many times do we run a program

# drop invalid external pipeline at the beginning
_invalid_external_pipeline_configs = []
for mode, arg_template in enabled_modes.items():
  if isinstance(arg_template, ExternalToolPipeline):
    if not arg_template.validate():
      print("External tool pipeline " + mode + " is not valid, dropping it")
      _invalid_external_pipeline_configs.append(mode)
for mode in _invalid_external_pipeline_configs:
  del enabled_modes[mode]

#-------------------------------------------------------------------------------
# compilation mode
#-------------------------------------------------------------------------------
debug_log_prefix = 'psaneval_'
debug_log_suffix = '_errorlog.txt'
makefile_outpath = 'Makefile'

compilation_lock = threading.Lock()

def print_output_file(outfile):
  now = datetime.datetime.now()
  compilation_lock.acquire()
  print('[' + now.strftime("%H:%M:%S") + "] " + outfile)
  compilation_lock.release()

def get_disaasemble_command(infile : str, outfile : str):
  if os.path.splitext(infile)[-1] == os.path.splitext(outfile)[-1]:
    return [ "cp", infile, outfile ]
  return [ llvmdis_path, infile, '-o', outfile]

def get_analyzer_command(args_template : typing.List[str], infile : str, outfile : str):
  args = [ psan_path ]
  for a in args_template:
    if a.startswith('{'):
      if a == "{in}":
        args.append(infile)
      elif a == "{out}":
        args.append(outfile)
      else:
        raise RuntimeError("Unknown special operand")
    else:
      args.append(a)
  return args

def get_analyzer_make_command(args_template : typing.List[str], infile : str, outfile : str):
  if isinstance(arg_template, ExternalToolPipeline):
    if arg_template.require_optimized_input:
      preopt_filename = os.path.join(os.path.dirname(outfile), os.path.basename(outfile) + "_preopt.ll")
      preopt_commands = get_preoptimize_command(infile, preopt_filename, False)
      actual_commands = arg_template.get_complete_commands(preopt_filename, outfile)
      return preopt_commands + ["&&"] + actual_commands
    return arg_template.get_complete_commands(infile, outfile)
  if len(args_template) == 0:
    return get_disaasemble_command(infile, outfile)
  return get_analyzer_command(args_template, infile, outfile)

def run_analyzer(args_template : typing.List[str], infile : str, outfile : str, is_export_stats : bool, debug_name: str) -> typing.Optional[dict]:
  print_output_file(outfile)
  # run the analyzer for the version, return the compile-time stats if available
  if len(args_template) == 0:
    # disassemble
    args = get_disaasemble_command(infile, outfile)
    result = subprocess.run(args, capture_output = True)
    if result.returncode != 0:
      raise RuntimeError("llvm-dis invocation failed: " + str(args))
    return None

  # invoke the analyzer
  # prepare the argument list first
  args = get_analyzer_command(args_template, infile, outfile)
  result = subprocess.run(args, capture_output = True)
  if result.returncode != 0:
    # only print last few lines of stdout
    stdout_str = result.stdout.decode()
    stderr_str = result.stderr.decode()
    stdout_lines = stdout_str.splitlines()
    compilation_lock.acquire()
    numlines = 20
    if (len(stdout_lines) > numlines):
      stdout_lines = stdout_lines[-numlines:]
    print('stdout:')
    print('\n'.join(stdout_lines))
    print('stderr:')
    print(stderr_str)
    # write complete log before crashing
    debug_log_path = debug_log_prefix + debug_name + debug_log_suffix
    with open(debug_log_path, 'w') as f:
      f.write("PSAN invocation failed: " + str(args))
      f.write("\nstderr:\n")
      f.write(stdout_str)
      f.write("\nstderr:\n")
      f.write(stderr_str)
      print('Error log written to ' + debug_log_path)
    compilation_lock.release()
    raise RuntimeError("PSAN invocation failed: " + str(args))
  if not is_export_stats:
    return None
  if not os.path.isfile(outfile):
    raise RuntimeError("PSAN is not creating output file: " + str(args))
  lines = result.stdout.split(b'\n')
  numlines = len(lines)
  for i in range(numlines-1, 0, -1): # it is safe to exclude line 0
    curline = lines[i]
    if curline.startswith("Pack:".encode()):
      result_line = lines[i+1]
      return json.loads(result_line)
  raise RuntimeError("PSAN invocation cannot extract compile-time stats")

def get_compiler_command(program_set_name : str, program_name: str, ir_file : str, output_file : str, additional_flags : list[str], is_cpp_program):
  command_list = compilation_command_info[0]
  if program_set_name in compilation_command_info[1]:
    curnode = compilation_command_info[1][program_set_name]
    if curnode[0] is not None:
      command_list = curnode[0]
    if program_name in curnode[1]:
      curnode = curnode[1][program_name]
      if curnode[0] is not None:
        command_list = curnode[0]
  compiler_path = clangpp_path if is_cpp_program else clang_path
  result_list = [compiler_path]
  for s in command_list:
    if s == "{in}":
      result_list.append(ir_file)
    elif s == "{out}":
      result_list.append(output_file)
    else:
      result_list.append(s)
  if additional_flags is not None:
    result_list.extend(additional_flags)
  return result_list

def get_preoptimize_command(ir_file : str, output_file : str, is_cpp_program : bool):
  # we need to optimize the input IR first
  # this is only for external tool pipeline
  compiler_path = clangpp_path if is_cpp_program else clang_path
  return [compiler_path, "-O3", "-S", "-emit-llvm", ir_file, "-o", output_file]

def run_compiler(program_set_name : str, program_name: str, ir_file : str, output_file : str, additional_flags : list[str], is_cpp_program : bool):
  print_output_file(output_file)
  result = subprocess.run(get_compiler_command(program_set_name, program_name, ir_file, output_file, additional_flags, is_cpp_program), capture_output = True, check=False)
  if len(result.stderr) > 0:
    compilation_lock.acquire()
    print(result.stderr.decode())
    compilation_lock.release()
  if result.returncode != 0:
    raise RuntimeError("Compilation failed for " + ir_file)

def time_wpa_svfg(input_path : str) -> int:
  args = timer_command_prefix_list + [psan_wpa, "-ander", "-svfg", input_path]
  result = subprocess.run(args, capture_output = True, check=False)
  stderr = result.stderr.decode()
  if result.returncode != 0:
    if len(stderr) > 0:
      compilation_lock.acquire()
      print(stderr)
      compilation_lock.release()
    raise RuntimeError("WPA SVFG failed for " + input_path)
  stderr_lines = stderr.splitlines()
  parse_result = parse_log_data(stderr_lines)
  if timestat := parse_result[0]:
    return int(timestat.time*1000) # convert to ms
  compilation_lock.acquire()
  print(stderr)
  compilation_lock.release()
  raise RuntimeError("WPA SVFG failed to extract time")

def mkdir_p(path : str):
  p = pathlib.Path(path)
  p.mkdir(parents=True, exist_ok=True)

def get_xlsx_data_str(data, headername) -> str | int:
  if isinstance(data, str):
    # if the data is in the form of "X/Y" where X and Y are both integers AND X > 0,
    # we also output the percentage (i.e., change "X/Y" to "X/Y (P%)")
    slash_splitted_parts = data.split('/')
    if len(slash_splitted_parts) == 2 and slash_splitted_parts[0].isdigit() and slash_splitted_parts[1].isdigit():
      x, y = int(slash_splitted_parts[0]), int(slash_splitted_parts[1])
      if x > 0:
        # Calculate the percentage
        percentage = (x / y) * 100
        # Return the formatted string
        return f"{x}/{y} ({percentage:.2f}%)"
    return data
  if headername.startswith("Time_"):
    # if we are recording time, the data value should be a number in ms
    # we want to convert it to hr/min/s/ms
    assert isinstance(data, int)
    hrs = 0
    mins = 0
    sec = 0
    ms = data
    result_str_list = []
    if ms >= 1000:
      sec = ms // 1000
      ms = ms % 1000
      if sec >= 60:
        mins = sec // 60
        sec = sec % 60
        if mins >= 60:
          hrs = mins // 60
          mins = mins % 60
          result_str_list.append(str(hrs) + 'h')
        result_str_list.append(str(mins) + 'm')
      result_str_list.append(str(sec) + 's')
    result_str_list.append(str(ms) + 'ms')
    return ' '.join(result_str_list) + '(' + str(data) + ')'
  if isinstance(data, int):
    return data
  return str(data)

compilation_stat_dict = {}
compilation_stat_lock = threading.Lock()

def handle_program_compilation(instrument : bool, compilation : bool, time_svfg : bool, program_set_name : str, program_name : str, mode : str, arg_template : typing.List[str] | ExternalToolPipeline, input_path : str, temp_dir : str, output_dir : str):
  basename = program_name + '_' + mode
  ir_name = basename + '.' + ir_suffix_temp
  temp_out_path = os.path.join(temp_dir, ir_name)
  output_path = os.path.join(output_dir, basename)
  if instrument:
    is_export_stats = (mode == static_stat_mode)
    if isinstance(arg_template, ExternalToolPipeline):
      ir_file_in = input_path
      if arg_template.require_optimized_input:
        preopt_filename = os.path.join(temp_dir, basename + "_preopt.ll")
        print_output_file(preopt_filename)
        preopt_commands = get_preoptimize_command(input_path, preopt_filename, False)
        result = subprocess.run(preopt_commands, capture_output = True, check=False)
        if result.returncode != 0:
          compilation_lock.acquire()
          print(result.stderr.decode())
          compilation_lock.release()
          raise RuntimeError("Pre-optimization failed: " + str(preopt_commands))
        ir_file_in = preopt_filename
      print_output_file(temp_out_path)
      cmd = arg_template.get_complete_commands(ir_file=ir_file_in, instrumented_ir_out=temp_out_path)
      result = subprocess.run(cmd, capture_output = True, check=False)
      if len(result.stderr) > 0:
        compilation_lock.acquire()
        print(result.stderr.decode())
        compilation_lock.release()
      if result.returncode != 0:
        raise RuntimeError("External tool pipeline invocation failed: " + str(cmd))
    else:
      result = run_analyzer(arg_template, input_path, temp_out_path, is_export_stats, basename)
      svfg_time = None
      if is_export_stats and time_svfg:
        svfg_time = time_wpa_svfg(input_path)
        result["Time_Extra_SVFG"] = svfg_time
      if is_export_stats:
        compilation_stat_lock.acquire()
        if program_set_name not in compilation_stat_dict:
          compilation_stat_dict[program_set_name] = {}
        compilation_stat_dict[program_set_name][program_name] = result
        compilation_stat_lock.release()
  if not os.path.isfile(temp_out_path):
    if instrument:
      raise RuntimeError("Failed to create IR: " + temp_out_path)
    else:
      raise RuntimeError("IR file not found (need to generate first): " + temp_out_path)
  # then the compiler
  if compilation:
    additional_flags = []
    if isinstance(arg_template, ExternalToolPipeline):
      additional_flags = arg_template.get_ldflags()
    run_compiler(program_set_name, program_name, temp_out_path, output_path, additional_flags, True if program_name in cpp_programs_list else False)
  return None

def run_compilation(instrument : bool = False, compilation : bool = False, time_svfg : bool = False,
                    numthreads : int | None = None, mode_filter : list[str] | None = None, progset_filter : list[str] | None = None):
  if progset_filter is not None and len(progset_filter) == 0:
    progset_filter = None
  if mode_filter is not None and len(mode_filter) == 0:
    mode_filter = None
  with concurrent.futures.ThreadPoolExecutor(max_workers=numthreads) as executor:
    future_list = []
    for program_set_name, program_list in enabled_programs_dict.items():
      if progset_filter is not None and program_set_name not in progset_filter:
        continue
      input_dir = os.path.join(src_ir_path, program_set_name)
      temp_dir = os.path.join(temp_ir_path, program_set_name)
      output_dir = os.path.join(root_output_path, program_set_name)
      mkdir_p(temp_dir)
      mkdir_p(output_dir)
      for p in program_list:
        src_name = p + '.' + ir_suffix_in
        input_path = os.path.join(input_dir, src_name)
        if not os.path.isfile(input_path):
          raise RuntimeError("Input file not found: " + input_path)
        stat_dict = None
        for mode, arg_template in enabled_modes.items():
          if mode_filter is not None and mode not in mode_filter:
            continue
          future_list.append(executor.submit(handle_program_compilation, instrument, compilation, time_svfg, program_set_name, p, mode, arg_template, input_path, temp_dir, output_dir))
    for future in concurrent.futures.as_completed(future_list):
      future.result()

  print_output_file("(completed)")

  # code below are only for instrumentation mode
  if not instrument:
    return

  # if we enabled any filter, we also skip the stats generation
  if mode_filter is not None or progset_filter is not None:
    return

  stat_additional_dict = {} # stat_name (e.g., "S1Taint_InitialCause") -> additional_info[key] (e.g., "ExternCall") -> value (set -> program -> dict[str, int])
  stat_basedata_dict = {} # set -> program -> dict[str, str]
  additional_stat_columns = []
  stat_additional_overview_dict = {} # stat_name -> key -> program -> int
  for program_set_name, program_list in enabled_programs_dict.items():
    set_basedata_dict = {}
    stat_basedata_dict[program_set_name] = set_basedata_dict
    for p in program_list:
      additional_stat_columns.append(p) # we do not include set name because if the columns are too narrow, the set names are occupying too much space
      stat_dict = compilation_stat_dict[program_set_name][p]
      if stat_dict is None:
        raise RuntimeError("no mode is providing stats")
      print("  stat for " + program_set_name + "/" + p + ": " + str(stat_dict))
      # convert the stat_dict into two parts:
      # 1. the part that is simply dict[str, str], write to stat_basedata_dict
      # 2. additional data that does to distinct files, directly write to stat_additional_dict
      raw_stat_dict = stat_dict
      program_basedata_dict = {}
      set_basedata_dict[p] = program_basedata_dict
      for k, v in raw_stat_dict.items():
        if isinstance(v, (str, int)):
          program_basedata_dict[k] = get_xlsx_data_str(v, k)
        elif isinstance(v, dict):
          if k not in stat_additional_overview_dict:
            stat_additional_overview_dict[k] = {}
          for k2, v2 in v.items():
            if k2 == "AdditionalInfo":
              if k not in stat_additional_dict:
                cur_stat_additional_dict = {}
                stat_additional_dict[k] = cur_stat_additional_dict
              else:
                cur_stat_additional_dict = stat_additional_dict[k]
              for k3, v3 in v2.items():
                # we expect v3 to be a dict[str, int]
                if not isinstance(v3, dict):
                  raise RuntimeError("Expected dict[str, int] for additional info")
                if k3 not in cur_stat_additional_dict:
                  cur_stat_value_dict = {}
                  cur_stat_additional_dict[k3] = cur_stat_value_dict
                else:
                  cur_stat_value_dict = cur_stat_additional_dict[k3]
                if program_set_name not in cur_stat_value_dict:
                  cur_stat_value_dict[program_set_name] = {}
                cur_stat_value_dict[program_set_name][p] = v3
            else:
              assert isinstance(v2, (str, int))
              composite_key = k + '_' + k2
              if k2 not in stat_additional_overview_dict[k]:
                stat_additional_overview_dict[k][k2] = {}
              stat_additional_overview_dict[k][k2][p] = get_xlsx_data_str(v2, composite_key)
        else:
          raise RuntimeError("Unknown stat type: " + str(v))
  # make stat_additional_dict complete
  stat_additional_row_headers = {}
  for stat_name, stat_values in stat_additional_dict.items():
    stat_additional_row_headers[stat_name] = {}
    for stat_value, cur_dict in stat_values.items():
      key_dict : dict[str, bool] = {}
      for set_name, set_details in cur_dict.items():
        for program_name, details in set_details.items():
          for k in details.keys():
            key_dict[k] = True
      key_list = list(key_dict.keys())
      key_list.sort()
      stat_additional_row_headers[stat_name][stat_value] = key_list
      for program_set_name, program_list in enabled_programs_dict.items():
        if program_set_name not in cur_dict:
          cur_dict[program_set_name] = {}
        cur_program_set_dict = cur_dict[program_set_name]
        for p in program_list:
          if p not in cur_program_set_dict:
            cur_program_set_dict[p] = {}
          cur_program_data_dict = cur_program_set_dict[p]
          for k in key_dict.keys():
            if k not in cur_program_data_dict:
              cur_program_data_dict[k] = 0

  # Create a new workbook
  wb = openpyxl.Workbook()

  # Add BaseData sheet
  base_data_sheet = wb.active
  base_data_sheet.title = "BaseData"

  # Write headers for BaseData sheet
  base_data_headers = ['Set', 'Program']
  for set_name, programs in stat_basedata_dict.items():
    for program_name, stats in programs.items():
      for stat in stats.keys():
        if stat not in base_data_headers:
          base_data_headers.append(stat)
  base_data_sheet.append(base_data_headers)

  column_textwidth_max = 40
  column_textwidth_min = 10

  # Write data for BaseData sheet
  base_data_text_maxlen = [max(column_textwidth_min, len(s)) for s in base_data_headers]
  for set_name, programs in stat_basedata_dict.items():
    if len(set_name) > base_data_text_maxlen[0]:
      base_data_text_maxlen[0] = len(set_name)
    for program_name, stats in programs.items():
      if len(program_name) > base_data_text_maxlen[1]:
        base_data_text_maxlen[1] = len(program_name)
      row = [set_name, program_name] + [stats.get(stat, '') for stat in base_data_headers[2:]]
      base_data_sheet.append(row)
      for i in range(2, len(base_data_headers)):
        cell_len = len(str(row[i]))
        if cell_len > base_data_text_maxlen[i]:
          base_data_text_maxlen[i] = cell_len
  for i in range(len(base_data_headers)):
    base_data_sheet.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width = min(base_data_text_maxlen[i], column_textwidth_max) + 2

  # Add sheets for additional stats
  for stat_name, stat_values in stat_additional_dict.items():
    overview_sheet = wb.create_sheet(title=stat_name)
    overview_sheet.append([stat_name] + additional_stat_columns)
    textlen = max(column_textwidth_min, len(stat_name))
    for key, details in stat_additional_overview_dict[stat_name].items():
      if len(key) > textlen:
        textlen = len(key)
      row = [key]
      for p in additional_stat_columns:
        row.append(details.get(p, None))
      overview_sheet.append(row)
    overview_sheet.column_dimensions['A'].width = min(textlen, column_textwidth_max) + 2
    for stat_value, programs in stat_values.items():
      sheet_name = stat_value # if we combine stat_name and stat_value, the sheet name will be too long
      additional_sheet = wb.create_sheet(title=sheet_name)

      # Write headers for additional stats sheet
      additional_headers = ['Value'] + additional_stat_columns
      additional_sheet.append(additional_headers)

      # Write data for additional stats sheet
      rows_list = stat_additional_row_headers[stat_name][stat_value]
      maxkeylen = column_textwidth_min
      for key in rows_list:
        if len(key) > maxkeylen:
          maxkeylen = len(key)
        row = [key]
        for program_set_name, program_list in enabled_programs_dict.items():
          for p in program_list:
            row_data = None
            if program_set_name in programs and p in programs[program_set_name]:
              row_data = programs[program_set_name][p].get(key, None)
              if row_data == 0:
                row_data = None
            row.append(row_data)
        additional_sheet.append(row)
      additional_sheet.column_dimensions['A'].width = min(maxkeylen, column_textwidth_max) + 2

  # Save the workbook
  wb.save(static_stats_output_path_xlsx)

def run_makefile_generation():
  body_part = ''
  defines = ''
  all_target_list = []
  clean_list = []

  defines += 'SRC_IR_PATH=' + src_ir_path + '\n'
  defines += 'TMP_IR_PATH=' + temp_ir_path + '\n'
  defines += 'OUTPUT_PATH=' + root_output_path + '\n'
  defines += '\n'

  for program_set_name, program_list in enabled_programs_dict.items():
    input_dir  = '$(SRC_IR_PATH)/' + program_set_name
    temp_dir   = '$(TMP_IR_PATH)/' + program_set_name
    output_dir = '$(OUTPUT_PATH)/' + program_set_name
    for p in program_list:
      src_name = p + '.' + ir_suffix_in
      input_path = os.path.join(input_dir, src_name)
      for mode, arg_template in enabled_modes.items():
        basename = p + '_' + mode
        ir_name = basename + '.' + ir_suffix_temp
        temp_out_path = os.path.join(temp_dir, ir_name)
        output_path = os.path.join(output_dir, basename)
        all_target_list.append(output_path)
        clean_list.append(temp_out_path)
        clean_list.append(output_path)
        # write rule for the analyzer run
        body_part += temp_out_path + ': ' + input_path + '\n'
        body_part += '\tmkdir -p ' + temp_dir + '\n'
        body_part += '\t@ ' + ' '.join(get_analyzer_make_command(arg_template, input_path, temp_out_path)) + '\n\n'
        # write rule for compiler run
        body_part += output_path + ': ' + temp_out_path + '\n'
        body_part += '\tmkdir -p ' + output_dir + '\n'
        body_part += '\t@ ' + ' '.join(get_compiler_command(
          program_set_name=program_set_name,
          program_name=p,
          ir_file=temp_out_path,
          output_file=output_path,
          additional_flags=arg_template.get_ldflags() if isinstance(arg_template, ExternalToolPipeline) else [],
          is_cpp_program=(True if p in cpp_programs_list else False))
        ) + '\n\n'
  with open(makefile_outpath, 'w', encoding='utf-8') as f:
    f.write(defines)
    all_target_str = 'all: ' + ' '.join(all_target_list) + '\n'
    f.write(all_target_str)
    f.write('.PHONY: clean\n\n')
    f.write(body_part)
    clean_str = 'clean:\n'
    clean_str += '\trm -f ' + ' '.join(clean_list) + '\n'
    f.write(clean_str)


#-------------------------------------------------------------------------------
# generation mode
#-------------------------------------------------------------------------------

def helper_get_run_args(program_set_name : str, program_name : str, datafile_path : str) -> str:
  key_str = program_set_name + '/' + program_name
  if key_str not in program_command_dict:
    raise RuntimeError("Command not specified for program: " + key_str)
  command = program_command_dict[key_str]
  # sanity check: if file_path is used, the directory should exist
  if '{file_path}' in command:
    if not os.path.isdir(datafile_path):
      raise RuntimeError("Data file path does not exist: " + datafile_path)
  return command.format(file_path=datafile_path)

def helper_add_redirection(entire_command : list[str], stdout_path : str):
  if any(['>' in s for s in entire_command]):
    entire_command.append('2>')
    entire_command.append(stdout_path)
  else:
    entire_command.append('>')
    entire_command.append(stdout_path)
    entire_command.append("2>&1")

def helper_form_execute_command(command_prefix : str, executable_path_from_script : str, args : str, datafile_path : str, run_script_base_path : str) -> list[str]:
  if not os.path.isdir(datafile_path):
    return [command_prefix, executable_path_from_script, args]
  # this command should use the datafile_path as the working directory (to take care of implicit file dependencies)
  executable_path_updated = os.path.abspath(os.path.join(run_script_base_path, executable_path_from_script))
  return ['(cd', os.path.abspath(datafile_path), ';', command_prefix, executable_path_updated, args, ')']

def get_run_command(program_set_name : str, program_name : str, executable_path : str, stdout_path : str, datafile_path : str, run_script_base_path : str, use_integrated_log : bool) -> str:
  args = helper_get_run_args(program_set_name, program_name, datafile_path)
  entire_command = helper_form_execute_command(timer_command_prefix, executable_path, args, datafile_path, run_script_base_path)
  if not use_integrated_log:
    helper_add_redirection(entire_command, stdout_path)
  return ' '.join(entire_command)

def get_perfstat_run_command(program_set_name : str, program_name : str, executable_path : str, stdout_path : str, datafile_path : str, run_script_base_path : str) -> str:
  args = helper_get_run_args(program_set_name, program_name, datafile_path)
  entire_command = helper_form_execute_command(perfstat_command_prefix, executable_path, args, datafile_path, run_script_base_path)
  helper_add_redirection(entire_command, stdout_path)
  return ' '.join(entire_command)

def get_run_echo(program_set_name : str, program_name : str, executable_path : str, runIndex : int, totalRun : int) -> str:
  return 'echo "Running {program_name} [{index}/{count}]..."'.format(program_name = executable_path, index = runIndex+1, count = totalRun)

def run_generation():
  # regression and split mode scripts are intended to run in QEMU, and the integrated version is intended for FPGA (where the file system is mounted read-only)
  # regression script should only run all versions (modes) of programs once
  # other evaluation ones (integrated/split) should run the modes listed in dynamic_time_stat_modes N times and run other versions once
  run_script_text_integrated  = "#!/bin/sh\n"
  run_script_text_split       = "#!/bin/sh\n"
  run_script_text_regression  = "#!/bin/sh\n"
  run_script_text_perfstat    = "#!/bin/sh\n"
  run_script_text_perfstat_1  = "#!/bin/sh\n"
  run_script_text_rt_stats    = "#!/bin/sh\n"
  run_script_text_split_fixms = "#!/bin/sh\n" # run only if output log not already exists (e.g., if test case fails and we have removed the output log)
  makefile_regression_targets = {}
  mkdir_p(run_script_base_path)

  def get_print_date_cmd(timename : str):
    return f'echo "{timename}: `date --rfc-3339=seconds`"\n'
  start_time_str = get_print_date_cmd("start")
  run_script_text_integrated += start_time_str
  run_script_text_split      += start_time_str
  run_script_text_regression += start_time_str
  run_script_text_perfstat   += start_time_str
  run_script_text_perfstat_1 += start_time_str
  run_script_text_rt_stats   += start_time_str
  run_script_text_split_fixms+= start_time_str

  logfile_dir = os.path.relpath(run_log_rootdir, run_script_base_path)
  datafile_dir = os.path.relpath(run_data_rootdir, run_script_base_path)
  perfstat_log_dir = os.path.relpath(run_perfstat_log_rootdir, run_script_base_path)

  regression_functions = []

  for program_set_name, program_list in enabled_programs_dict.items():
    output_dir = os.path.join(root_output_path, program_set_name)
    datafile_prog_set_path = os.path.join(datafile_dir, program_set_name)
    filepath_overrides = None
    if program_set_name in program_filepath_overrides:
      filepath_overrides = program_filepath_overrides[program_set_name]
    if isinstance(filepath_overrides, str):
      datafile_prog_set_path = filepath_overrides
    for p in program_list:
      if program_set_name + p in compile_only_programs:
        continue
      datafile_path = os.path.join(datafile_prog_set_path, p)
      if isinstance(filepath_overrides, dict) and p in filepath_overrides:
        datafile_path = filepath_overrides[p]
      makefile_regression_program_target = "regr_" + program_set_name + "_" + p
      makefile_regression_program_target_dict = {}
      makefile_regression_targets[makefile_regression_program_target] = makefile_regression_program_target_dict
      regression_function_name = "runpprogram_" + program_set_name + "_" + p
      regression_function_name = regression_function_name.replace('-', '_')
      regression_function_name = regression_function_name.replace('.', '_')
      run_script_text_regression += regression_function_name + '() {\n'
      regression_functions.append(regression_function_name)
      for mode in enabled_modes.keys():
        basename = p + '_' + mode
        output_path = os.path.join(output_dir, basename)
        # prepare the run command
        executable_relative_path = os.path.relpath(output_path, run_script_base_path)
        is_nontiming_mode = mode not in dynamic_time_stat_modes
        is_rt_stats_mode = mode in dynamic_custom_stat_modes

        for i in range(0, num_repeat_execution):
          logfilename = "_".join([program_set_name, p, mode, str(i)]) + ".txt"
          stdout_path = os.path.join(logfile_dir, logfilename)
          perfstat_logpath = os.path.join(perfstat_log_dir, logfilename)
          echo_command = get_run_echo(program_set_name, p, executable_relative_path, i, num_repeat_execution)

          # write the echo to the scripts
          if not is_nontiming_mode:
            run_script_text_integrated  += echo_command + '\n'
            run_script_text_split       += echo_command + '\n'
            run_script_text_perfstat    += echo_command + '\n'
          if i == 0:
            run_script_text_regression  += echo_command + '\n'
            if not is_nontiming_mode:
              run_script_text_perfstat_1  += echo_command + '\n'
            if is_rt_stats_mode:
              run_script_text_rt_stats += echo_command + '\n'

          # write the actual commands
          run_cmd_normallog     = get_run_command(program_set_name, p, executable_relative_path, stdout_path, datafile_path, run_script_base_path, False)
          run_cmd_integratedlog = get_run_command(program_set_name, p, executable_relative_path, stdout_path, datafile_path, run_script_base_path, True)
          run_cmd_perfstat      = get_perfstat_run_command(program_set_name, p, executable_relative_path, perfstat_logpath, datafile_path, run_script_base_path)
          if not is_nontiming_mode:
            run_script_text_integrated += "echo '" + run_script_integrated_log_start + ' '.join([program_set_name, p, mode, str(i)]) + "'\n"
            run_script_text_integrated += run_cmd_integratedlog + '\n'
            run_script_text_split      += run_cmd_normallog + '\n'
            run_script_text_perfstat   += run_cmd_perfstat + '\n'
            run_script_text_integrated += "echo '" + run_script_integrated_log_end + ' '.join([program_set_name, p, mode, str(i)]) + "'\n"
            # for run_script_text_split_fixms, we first test if the log file exists
            # and if it does not exist, we run the command
            run_script_text_split_fixms += 'if [ ! -f ' + stdout_path + ' ]; then\n'
            run_script_text_split_fixms += '\t' + echo_command + '\n'
            run_script_text_split_fixms += '\t' + run_cmd_normallog + '\n'
            run_script_text_split_fixms += 'fi\n'
          if i == 0:
            run_script_text_regression  += run_cmd_normallog + '\n'
            if mode != "baseline":
              makefile_regression_target = "regr_" + program_set_name + "_" + p + "_" + mode
              makefile_regression_program_target_dict[makefile_regression_target] = run_cmd_normallog
            if not is_nontiming_mode:
              run_script_text_perfstat_1  += run_cmd_perfstat + '\n'
            if is_rt_stats_mode:
              run_script_text_rt_stats += run_cmd_normallog + '\n'
      run_script_text_regression += '}\n'

  for regression_function_name in regression_functions:
    run_script_text_regression += regression_function_name + ' &\n'
  run_script_text_regression += 'wait\n'
  finish_time_str = get_print_date_cmd("finish")
  run_script_text_integrated += finish_time_str
  run_script_text_split      += finish_time_str
  run_script_text_regression += finish_time_str
  run_script_text_perfstat   += finish_time_str
  run_script_text_perfstat_1 += finish_time_str
  run_script_text_rt_stats   += finish_time_str
  run_script_text_split_fixms+= finish_time_str

  # write run scripts
  with open(os.path.join(run_script_base_path, run_script_name_integrated), 'w', encoding="utf-8") as f:
    f.write(run_script_text_integrated)
  with open(os.path.join(run_script_base_path, run_script_name_split),      'w', encoding="utf-8") as f:
    f.write(run_script_text_split)
  with open(os.path.join(run_script_base_path, run_script_name_regression), 'w', encoding="utf-8") as f:
    f.write(run_script_text_regression)
  with open(os.path.join(run_script_base_path, run_script_name_perfstat), 'w', encoding="utf-8") as f:
    f.write(run_script_text_perfstat)
  with open(os.path.join(run_script_base_path, run_script_name_perfstat_1), 'w', encoding="utf-8") as f:
    f.write(run_script_text_perfstat_1)
  with open(os.path.join(run_script_base_path, run_script_name_rt_stats), 'w', encoding="utf-8") as f:
    f.write(run_script_text_rt_stats)
  with open(os.path.join(run_script_base_path, run_script_name_split_fixms), 'w', encoding="utf-8") as f:
    f.write(run_script_text_split_fixms)
  with open(os.path.join(run_script_base_path, makefile_regression_outpath), 'w', encoding="utf-8") as f:
    all_targets = ' \\\n\t'.join(makefile_regression_targets.keys())
    f.write(f'all: \\\n\t{all_targets}\n\n')
    for prog, prog_dict in makefile_regression_targets.items():
      f.write(prog + ':\n')
      for target, command in prog_dict.items():
        f.write('\t' + command + '\n')
    #for prog, prog_dict in makefile_regression_targets.items():
    #  for target, command in prog_dict.items():
    #    f.write(target + ':\n')
    #    f.write('\t' + command + '\n')

  # done

#-------------------------------------------------------------------------------
# parsing mode
#-------------------------------------------------------------------------------

@dataclasses.dataclass
class TimeStat:
  time : float # real time in seconds
  mem : int # maximum resident set size, in KB

psan_packed_stat_head = 'PSan packed stat: '
#time_time_stat_regex = '(>\s+)?(?P<time>\d*\.\d+) real\s+(?P<user_time>\d*\.\d+) user\s+(?P<sys_time>\d*\.\d+) sys'
#time_mem_stat_end = 'maximum resident set size'
time_time_stat_prefix = 'Elapsed (wall clock) time (h:mm:ss or m:ss): '
# matches h:mm:ss or m:ss
time_time_stat_text_regex = '(?P<hour>\d+):(?P<minute>\d+):(?P<second>[\d.]+)|(?P<minute2>\d+):(?P<second2>[\d.]+)'
time_mem_stat_prefix = 'Maximum resident set size (kbytes):'
def parse_log_data(lines : typing.List[str]) -> typing.Tuple[typing.Optional[TimeStat], typing.Optional[dict[str, int]]]:
  time_time_stat = None
  time_mem_stat = None
  psan_stat = None
  for line in lines:
    content = line.strip()
    if content.startswith(psan_packed_stat_head):
      psan_stat_text = content[len(psan_packed_stat_head):].strip()
      psan_stat = json.loads(psan_stat_text)
    elif content.startswith(time_mem_stat_prefix):
      time_mem_text = content[len(time_mem_stat_prefix):].strip()
      assert time_mem_stat is None
      time_mem_stat = int(time_mem_text)
    elif content.startswith(time_time_stat_prefix):
      time_text = content[len(time_time_stat_prefix):].strip()
      assert time_time_stat is None
      match = re.match(time_time_stat_text_regex, time_text)
      if match is not None:
        hour = match.group('hour')
        minute = match.group('minute')
        second = match.group('second')
        minute2 = match.group('minute2')
        second2 = match.group('second2')
        assert (hour is not None and minute is not None and second is not None) or (minute2 is not None and second2 is not None)
        if hour is not None:
          time_time_stat = int(hour)*3600 + int(minute)*60 + float(second)
        else:
          time_time_stat = int(minute2)*60 + float(second2)
  time_stat = None
  if time_time_stat is not None:
    assert time_mem_stat is not None
    time_stat = TimeStat(time=time_time_stat, mem=time_mem_stat)
  return (time_stat, psan_stat)

def write_runtime_stats(time_stat_dict, psan_stat_dict, max_index) -> None:
  # time_stat_dict: [set/program] -> [mode] -> [index] -> TimeStat
  # psan_stat_dict: [mode] -> [set/program] -> [index] -> dict[str, int]

  # digested_time_stat_dict: [set/program] -> [mode] -> <avg_time, avg_mem> (both float)
  digested_time_stat_dict = {}

  # time stats
  time_stat_text = 'program,mode,avg_time,avg_mem'
  appeared_modes = set()
  for i in range(0, max_index+1):
    time_stat_text += ',time' + str(i) + ',mem' + str(i)
  time_stat_text += '\n'
  for program, program_data in sorted(time_stat_dict.items()):
    digested_time_stat_dict[program] = {}
    for mode, mode_data in sorted(program_data.items()):
      if mode not in appeared_modes:
        appeared_modes.add(mode)
      time_list = []
      mem_list = []
      for i in range(0, max_index+1):
        time_list.append(0.0)
        mem_list.append(0)
      num = 0
      sum_time = 0.0
      sum_mem = 0
      for run_index, stat in mode_data.items():
        time_list[run_index] = stat.time
        mem_list[run_index] = stat.mem
        num += 1
        sum_time += stat.time
        sum_mem += stat.mem
      avg_time = 0.0
      avg_mem = 0
      if num > 0:
        avg_time = sum_time / num
        avg_mem = float(sum_mem) / num
      digested_time_stat_dict[program][mode] = (avg_time, avg_mem)
      time_stat_text += program + ',' + mode + ',' + str(avg_time) + ',' + str(avg_mem)
      for i in range(0, max_index+1):
        time_stat_text += ',' + str(time_list[i]) + ',' + str(mem_list[i])
      time_stat_text += '\n'
  with open(runtime_time_stats_output_path, 'w', encoding="utf-8") as f:
    f.write(time_stat_text)

  # digested time stats (first half: summary_table_width; second half: ratio_table_width)
  #                | time(s)                | mem(KB)
  # <program name> | baseline <other modes> | baseline <other modes>
  # <------------------------ data ------------------------>
  # <space>
  # time_ratio     | <other modes>   | <space> | mem_ratio      | <other_modes>
  # <program name> | <ratio data>    |         | <program name> | <ratio data>

  appeared_dynamic_time_stat_modes = []
  for mode in dynamic_time_stat_modes:
    if mode in appeared_modes:
      appeared_dynamic_time_stat_modes.append(mode)
  # initialize the table headers
  max_table_width = len(appeared_dynamic_time_stat_modes)*2+1
  timestat_summary_table = []
  timestat_summary_table.append(['']*max_table_width)
  timestat_summary_table[0][1] = 'time(s)'
  timestat_summary_table[0][1+len(appeared_dynamic_time_stat_modes)] = 'mem(KB)'
  timestat_summary_table.append(['']*max_table_width)

  timestat_ratio_table = []
  timestat_ratio_table.append(['']*max_table_width)
  timestat_ratio_table[0][0] = 'time_ratio'
  timestat_ratio_table[0][1+len(appeared_dynamic_time_stat_modes)] = 'mem_ratio'
  mode_index = 0
  for mode in appeared_dynamic_time_stat_modes:
    # header for summary table
    timestat_summary_table[1][1+mode_index] = mode
    timestat_summary_table[1][1+len(appeared_dynamic_time_stat_modes)+mode_index] = mode
    # header for ratio table
    # exclude the baseline
    if mode_index > 0:
      timestat_ratio_table[0][mode_index] = mode
      timestat_ratio_table[0][1+len(appeared_dynamic_time_stat_modes)+mode_index] = mode
    mode_index += 1
    continue

  # write the data rows
  for program, program_data in sorted(digested_time_stat_dict.items()):
    # (these two must use different list objects)
    cur_summary_row = ['']*max_table_width
    cur_ratio_row = ['']*max_table_width
    # for ratio table, we want to exclude the program set name
    program_basename = program.split('/')[-1]
    assert isinstance(program_basename, str)
    cur_summary_row[0] = program
    cur_ratio_row[0] = program_basename
    cur_ratio_row[1+len(appeared_dynamic_time_stat_modes)] = program_basename

    baseline_avgtime = None
    baseline_avgmem = None
    mode_index = 0
    for mode in appeared_dynamic_time_stat_modes:
      avg_time, avg_mem = digested_time_stat_dict[program][mode]
      # fill the summary table
      cur_summary_row[1+mode_index] = str(avg_time)
      cur_summary_row[1+len(appeared_dynamic_time_stat_modes)+mode_index] = str(avg_mem)
      # fill the ratio table
      if baseline_avgtime is None:
        assert baseline_avgmem is None
        assert mode_index == 0
        baseline_avgtime = avg_time
        baseline_avgmem = avg_mem
      else:
        cur_ratio_row[mode_index] = str(avg_time/baseline_avgtime if baseline_avgtime > 0 else 0.0)
        cur_ratio_row[1+len(appeared_dynamic_time_stat_modes)+mode_index] = str(avg_mem/baseline_avgmem if baseline_avgmem > 0 else 0.0)
      mode_index += 1
    timestat_summary_table.append(cur_summary_row)
    timestat_ratio_table.append(cur_ratio_row)

  with open(runtime_time_stats_digested_output_path, 'w', encoding="utf-8") as f:
    for line in timestat_summary_table:
      f.write(','.join(line))
      f.write('\n')
    f.write(','.join(['']*max_table_width) + '\n')
    for line in timestat_ratio_table:
      f.write(','.join(line))
      f.write('\n')

  # psan stats
  # first, find all custom stat names
  psan_custom_stat_head = []
  for mode, mode_data in sorted(psan_stat_dict.items()):
    for program, program_data in sorted(mode_data.items()):
      for run, run_data in program_data.items():
        for k in run_data.keys():
          if k not in psan_custom_stat_head:
            psan_custom_stat_head.append(k)
  psan_stat_text = 'mode,program,' + ','.join(psan_custom_stat_head) + '\n'
  for mode, mode_data in sorted(psan_stat_dict.items()):
    for program, program_data in sorted(mode_data.items()):
      cur_data_dict = {}
      cur_data_numappear = {}
      for k in psan_custom_stat_head:
        cur_data_dict[k] = 0
        cur_data_numappear[k] = 0
      for run_index, stat in program_data.items():
        for k, v in stat.items():
          cur_data_dict[k] += v
          cur_data_numappear[k] += 1
      for k in psan_custom_stat_head:
        cur_data_dict[k] /= cur_data_numappear[k] if cur_data_numappear[k] > 0 else 1
      psan_stat_text += mode + ',' + program
      for k in psan_custom_stat_head:
        psan_stat_text += ',' + str(cur_data_dict[k])
      psan_stat_text += '\n'
  with open(runtime_psan_stats_output_path, 'w', encoding="utf-8") as f:
    f.write(psan_stat_text)

def populate_runtime_stats_to_dict(time_stat_dict, psan_stat_dict, set_name, program_name, mode_name, run_index, time_stat, psan_stat) -> None:
  entry_name = set_name + '/' + program_name
  if time_stat is not None:
    if mode_name in dynamic_time_stat_modes:
      if entry_name not in time_stat_dict:
        time_stat_dict[entry_name] = {}
      if mode_name not in time_stat_dict[entry_name]:
        time_stat_dict[entry_name][mode_name] = {}
      time_stat_dict[entry_name][mode_name][run_index]=time_stat
  if psan_stat is not None:
    if mode_name in dynamic_custom_stat_modes:
      if mode_name not in psan_stat_dict:
        psan_stat_dict[mode_name] = {}
      if entry_name not in psan_stat_dict[mode_name]:
        psan_stat_dict[mode_name][entry_name] = {}
      psan_stat_dict[mode_name][entry_name][run_index]=psan_stat

def parse_log_from_dir(time_stat_dict, psan_stat_dict, logdir : str) -> int:
  # time_stat_dict: [set/program] -> [mode] -> [index] -> TimeStat
  # psan_stat_dict: [mode] -> [set/program] -> [index] -> dict[str, int]
  logs = os.listdir(logdir)
  max_index = 0
  for log in logs:
    assert log.endswith('.txt')
    fields = log[:-4].split('_')
    assert len(fields) >= 4
    set_name = fields[0]
    program_name = '_'.join(fields[1:-2])
    mode_name = fields[-2]
    run_index = int(fields[-1])
    if run_index > max_index:
      max_index = run_index

    filepath = os.path.join(logdir, log)
    with open(filepath, 'r', encoding="utf-8") as f:
      lines = f.readlines()
      print('Parsing ' + filepath + ' ...')
      time_stat, psan_stat = parse_log_data(lines)
      populate_runtime_stats_to_dict(time_stat_dict, psan_stat_dict, set_name, program_name, mode_name, run_index, time_stat, psan_stat)
  return max_index

def parse_console_log(time_stat_dict, psan_stat_dict, logpath : str) ->  int:
  max_index = 0
  cur_lines = None # None: not in any part; list: inside a log
  set_name = ''
  program_name = ''
  mode_name = ''
  run_index = 0
  with open(logpath, 'r') as f:
    for line in f:
      if line.startswith(run_script_integrated_log_start):
        assert cur_lines is None
        cur_lines = []
        fields = line[len(run_script_integrated_log_start):].split(' ')
        assert len(fields) == 4
        set_name = fields[0]
        program_name = fields[1]
        mode_name = fields[2]
        run_index = int(fields[3])
        if run_index > max_index:
          max_index = run_index
        continue
      if line.startswith(run_script_integrated_log_end):
        assert isinstance(cur_lines, list)
        print('Parsing ' + set_name + '/' + program_name + '/' + mode_name + '-' + str(run_index) + ' ...')
        time_stat, psan_stat = parse_log_data(cur_lines)
        cur_lines = None
        populate_runtime_stats_to_dict(time_stat_dict, psan_stat_dict, set_name, program_name, mode_name, run_index, time_stat, psan_stat)
        continue
      # normal line handling
      # if we are inside a log section, add it; otherwise skip it
      if isinstance(cur_lines, list):
        cur_lines.append(line)
  # done
  return max_index


def run_log_parser(logpath : str) -> None:
  time_stat_dict = {} # [set/program] -> [mode] -> [index] -> TimeStat
  psan_stat_dict = {} # [mode] -> [set/program] -> [index] -> dict[str, int]
  if os.path.isdir(logpath):
    max_index = parse_log_from_dir(time_stat_dict, psan_stat_dict, logpath)
  elif os.path.isfile(logpath):
    max_index = parse_console_log(time_stat_dict, psan_stat_dict, logpath)
  elif os.path.exists(logpath):
    raise RuntimeError('Content exist but is neither directory or file: ' + logpath)
  else:
    raise RuntimeError('Log not found: ' + logpath)
  write_runtime_stats(time_stat_dict, psan_stat_dict, max_index)

#-------------------------------------------------------------------------------
# perf stat log parsing
#-------------------------------------------------------------------------------

def parse_perf_stat(output) -> dict[str, typing.Any]:
  result = {}
  # Start parsing after "Performance counter stats for ..."
  if isinstance(output, str):
    lines = output.split('\n')
  elif isinstance(output, list):
    lines = output
  else:
    raise RuntimeError("Unexpected output log type: " + str(type(output)))
  start_parsing = False
  for line in lines:
    if 'Performance counter stats for' in line:
      start_parsing = True
      continue
    if start_parsing:
      # Check for time elapsed and exit if found
      time_match = re.search(r'(\d+\.\d+) seconds time elapsed', line)
      if time_match:
        result['time-elapsed'] = float(time_match.group(1))
        break

      # Regular data lines
      data_match = re.match(r'\s*(\d+|\s*<not supported>\s*)\s+(\w[\w-]*)', line)
      if data_match:
        value, key = data_match.groups()
        # Convert "<not supported>" to None
        if value.strip() == '<not supported>':
          result[key] = None
        else:
          # Convert number strings to integers
          result[key] = int(value.replace(',', ''))
  return result

def parse_perf_stat_from_dir(logdir : str) -> int:
  if not os.path.isdir(logdir):
    raise RuntimeError('Log directory not found: ' + logdir)
  stat_dict = {} # [set/program] -> [mode] -> [index] -> data
  logs = os.listdir(logdir)
  max_index = 0
  keys_ordered = collections.OrderedDict()
  for log in logs:
    assert log.endswith('.txt')
    fields = log[:-4].split('_')
    assert len(fields) >= 4
    set_name = fields[0]
    program_name = '_'.join(fields[1:-2])
    mode_name = fields[-2]
    run_index = int(fields[-1])
    if run_index > max_index:
      max_index = run_index

    filepath = os.path.join(logdir, log)
    with open(filepath, 'r', encoding="utf-8") as f:
      lines = f.readlines()
      print('Parsing ' + filepath + ' ...')
      stat = parse_perf_stat(lines)
      for k in stat.keys():
        keys_ordered[k] = True
      if set_name not in stat_dict:
        stat_dict[set_name] = {}
      if program_name not in stat_dict[set_name]:
        stat_dict[set_name][program_name] = {}
      if mode_name not in stat_dict[set_name][program_name]:
        stat_dict[set_name][program_name][mode_name] = {}
      stat_dict[set_name][program_name][mode_name][run_index] = stat
  key_list = list(keys_ordered.keys())
  output_csv = []
  header = ','.join(["set", "program", "mode", "index"] + key_list) + "\n"
  output_csv.append(header)
  for program_set_name, program_set_details in enabled_programs_dict.items():
    if program_set_name in stat_dict:
      set_data = stat_dict[program_set_name]
      for program_name in program_set_details:
        if program_name in set_data:
          program_data = set_data[program_name]
          for mode in dynamic_time_stat_modes:
            if mode in program_data:
              mode_data = program_data[mode]
              for i in range(0, max_index+1):
                if i in mode_data:
                  run_data = mode_data[i]
                  row = [program_set_name, program_name, mode, str(i)]
                  for k in key_list:
                    row.append(str(run_data.get(k, '')))
                  output_csv.append(','.join(row) + "\n")
  output_csv.append("\n")
  with open(runtime_perf_stats_output_path, "w", encoding="utf-8") as f:
    f.writelines(output_csv)
  return max_index

#-------------------------------------------------------------------------------
# main
#-------------------------------------------------------------------------------

def main():
  parser = argparse.ArgumentParser(description='PSAN evaluation helper script')
  parser.add_argument('-n', action='store_true', help='run instrumentation on benchmarks')
  parser.add_argument('-c', action='store_true', help='compile benchmark programs')
  parser.add_argument('--time-svfg', action='store_true', help='time SVFG construction on source programs')
  parser.add_argument('-g', action='store_true', help='generate run scripts')
  parser.add_argument('-p', action='store_true', help='parse result log files')
  parser.add_argument('--perfstat', action='store_true', help='Parse perf stat logs instead of normal logs')
  parser.add_argument('-m', action='store_true', help='run makefile generation')
  parser.add_argument('--logpath', type=pathlib.Path, default=None, help='log file path / directory')
  parser.add_argument('--numthreads', type=int, default=None, help='number of threads to use (only for -n and -c)')
  parser.add_argument('--mode-filter', type=str, default=None, help='mode filter (only for -n and -c)')
  parser.add_argument('--progset-filter', type=str, default=None, help='program set filter (only for -n and -c)')
  args = parser.parse_args()
  if args.p and (args.n or args.c or args.g or args.m):
    parser.print_usage()
    parser.exit(status=1, message='Must be invoked with either -p or other modes, not both/all')
  if args.n or args.c:
    mode_filter = []
    if args.mode_filter is not None:
      mode_filter = [v for v in args.mode_filter.split(',') if len(v) > 0]
      for m in mode_filter:
        if m not in enabled_modes:
          raise RuntimeError('Invalid mode filter: ' + m)
    progset_filter = []
    if args.progset_filter is not None:
      progset_filter = [v for v in args.progset_filter.split(',') if len(v) > 0]
      for p in progset_filter:
        if p not in enabled_programs_dict:
          raise RuntimeError('Invalid program set filter: ' + p)
    run_compilation(instrument=args.n, compilation=args.c, time_svfg=args.time_svfg, numthreads=args.numthreads, mode_filter=mode_filter, progset_filter=progset_filter)
  if args.g:
    run_generation()
  if args.m:
    run_makefile_generation()
  if args.p:
    logpath = run_log_rootdir if not args.perfstat else run_perfstat_log_rootdir
    if args.logpath is not None:
      logpath = args.logpath.as_posix()
    if args.perfstat:
      parse_perf_stat_from_dir(logpath)
    else:
      run_log_parser(logpath)

if __name__ == "__main__":
  main()
