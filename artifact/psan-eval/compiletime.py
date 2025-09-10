#!/usr/bin/env python3

# this script is used to compile a report of analysis+instrumentation time from multiple PSAN reports
# the input xlsx files should be generated with the command `python3 ./psaneval.py -n --time-svfg`

import sys
import openpyxl
from collections import defaultdict

def parse_time_in_ms(time_str):
  """
  Parse a cell string like '3s 930ms(3930)' or '83ms(83)' or '2s 290ms(2290)'
  and return the integer inside parentheses as an integer (milliseconds).
  Returns None if parsing fails.
  """
  if not time_str:
    return None
  # We expect the time in parentheses at the end: e.g. '(3930)'
  # A simple approach: find the last '(' and parse until the closing ')'
  start_idx = time_str.rfind('(')
  if start_idx == -1:
    return None
  end_idx = time_str.rfind(')')
  if end_idx == -1 or end_idx <= start_idx:
    return None

  # Extract the substring
  ms_str = time_str[start_idx+1:end_idx].strip()
  try:
    return int(ms_str)
  except ValueError:
    return None

def main():
  if len(sys.argv) < 3:
    print("Usage: compiletime.py <output.xlsx> <input1.xlsx> [<input2.xlsx> ...]")
    sys.exit(1)

  output_xlsx = sys.argv[1]
  input_xlsx_files = sys.argv[2:]

  # We will store data in the following nested structure:
  # data[(setName, programName)][timeColumnName] = list of int (parsed ms)
  data = defaultdict(lambda: defaultdict(list))

  # We also want to record the set of all time columns we encounter
  all_time_columns = set()

  for filename in input_xlsx_files:
    print(f"Reading file: {filename}")
    wb = openpyxl.load_workbook(filename=filename, data_only=True)
    sheet = wb.worksheets[0]  # read only the first sheet

    # Find headers
    header_row = None
    for i, row in enumerate(sheet.iter_rows(values_only=True), start=1):
      if i == 1:
        # This is our header row
        header_row = list(row)
        break

    if not header_row:
      continue  # No rows in this sheet

    # Build a map from column name -> index
    col_index_map = {}
    for idx, col_name in enumerate(header_row):
      if col_name is not None:
        col_index_map[col_name] = idx

    # We need at least "Set" and "Program"
    if "Set" not in col_index_map or "Program" not in col_index_map:
      # Skip if missing these
      print(f"Warning: File {filename} does not contain 'Set' or 'Program' columns in the first sheet.")
      continue

    # Identify time columns (start with "Time_")
    time_columns = [c for c in col_index_map if c.startswith("Time_")]

    # Read data rows
    for row in sheet.iter_rows(min_row=2, values_only=True):
      # read the set and program
      set_val = row[col_index_map["Set"]]
      prog_val = row[col_index_map["Program"]]
      if set_val is None or prog_val is None:
        continue  # skip empty row

      # For each time column, parse
      for tcol in time_columns:
        idx = col_index_map[tcol]
        cell_val = row[idx]
        if cell_val is None:
          continue
        ms_val = parse_time_in_ms(str(cell_val))
        if ms_val is not None:
          data[(set_val, prog_val)][tcol].append(ms_val)
          all_time_columns.add(tcol)

  # Sort the time columns in a consistent order
  # You can reorder them as you prefer. We'll just sort alphabetically for output.
  all_time_columns = sorted(all_time_columns)

  # Now we compute the average for each (Set, Program, Time_Column)
  # We'll create a new workbook for output
  out_wb = openpyxl.Workbook()

  # 1) Sheet "Average"
  sheet_avg = out_wb.active
  sheet_avg.title = "Average"

  # Write the header
  avg_header = ["Set", "Program"] + all_time_columns
  sheet_avg.append(avg_header)

  # We'll sort rows by Set and Program (alphabetically).
  for (set_val, prog_val) in sorted(data.keys()):
    row_data = [set_val, prog_val]
    for tcol in all_time_columns:
      values = data[(set_val, prog_val)].get(tcol, [])
      if values:
        avg_ms = sum(values) / len(values)
      else:
        avg_ms = 0
      row_data.append(avg_ms)
    sheet_avg.append(row_data)

  # 2) Sheet "Summary"
  # We want columns: "Set", "Program", "Load+SVFG", "Analysis", "Instrumentation", "WPA"
  # Where:
  #    Load+SVFG = Time_IRLoad + Time_S3_SVF
  #    Analysis = sum of Time_S1_BeforeSolve, Time_S1_Solve, Time_S1_Dump, Time_S2_ClonePartition,
  #               Time_S2_Prune, Time_S2_Dump, Time_S3_InitialConstruction, Time_S3_Finish,
  #               Time_S4_Analysis, Time_S4_TypeSolving
  #    Instrumentation = sum of Time_S5_Preparation, Time_S5_Transform, Time_S5_MDProp, Time_S5_Wrapup
  #    WPA = Time_Extra_SVFG
  #
  # We'll do the "average of sums" approach, i.e., for each run:
  #   - load_plus_svfg_run_i = (Time_IRLoad_run_i + Time_S3_SVF_run_i)
  # Then we average across runs. That means we have to re-walk the raw data list.
  #
  # Alternatively, if there's guaranteed only one row per set/program per file, summing the
  # average time is effectively the same. But let's do it carefully to handle multiple runs.

  sheet_summary = out_wb.create_sheet("Summary")

  summary_header = ["Set", "Program", "Load+SVFG", "Analysis", "Instrumentation", "WPA"]
  sheet_summary.append(summary_header)

  # Define groups
  group_load_svfg = ["Time_IRLoad", "Time_S3_SVF"]
  group_analysis = [
    "Time_S1_BeforeSolve", "Time_S1_Solve", "Time_S1_Dump",
    "Time_S2_ClonePartition", "Time_S2_Prune", "Time_S2_Dump",
    "Time_S3_InitialConstruction", "Time_S3_Finish",
    "Time_S4_Analysis", "Time_S4_TypeSolving",
  ]
  group_instrumentation = [
    "Time_S5_Preparation", "Time_S5_Transform",
    "Time_S5_MDProp", "Time_S5_Wrapup",
  ]
  group_wpa = ["Time_Extra_SVFG"]

  # Helper to sum up times for a single run i
  def sum_times_for_run(timing_lists, run_index, time_cols):
    """Sum the run_index-th element across the given columns if it exists."""
    s = 0
    for c in time_cols:
      if c in timing_lists and len(timing_lists[c]) > run_index:
        s += timing_lists[c][run_index]
    return s

  # For each set/program, we might have multiple runs recorded in data[(set_val,prog_val)].
  # We want to combine them properly.
  for (set_val, prog_val) in sorted(data.keys()):
    timing_lists = data[(set_val, prog_val)]

    # The number of runs for this (Set,Program) is the maximum length among the lists
    # (in typical data, all time columns have the same number of runs, but let's be safe).
    num_runs = max(len(lst) for lst in timing_lists.values())

    load_svfg_vals = []
    analysis_vals = []
    instrumentation_vals = []
    wpa_vals = []

    for run_i in range(num_runs):
      load_svfg_vals.append(sum_times_for_run(timing_lists, run_i, group_load_svfg))
      analysis_vals.append(sum_times_for_run(timing_lists, run_i, group_analysis))
      instrumentation_vals.append(sum_times_for_run(timing_lists, run_i, group_instrumentation))
      wpa_vals.append(sum_times_for_run(timing_lists, run_i, group_wpa))

    # Compute average
    def avg_if_any(lst):
      return sum(lst)/len(lst) if lst else 0

    avg_load_svfg = avg_if_any(load_svfg_vals)
    avg_analysis = avg_if_any(analysis_vals)
    avg_instr = avg_if_any(instrumentation_vals)
    avg_wpa = avg_if_any(wpa_vals)

    sheet_summary.append([
      set_val,
      prog_val,
      avg_load_svfg,
      avg_analysis,
      avg_instr,
      avg_wpa
    ])

  # Finally, save the workbook
  out_wb.save(output_xlsx)
  print(f"Combined report written to {output_xlsx}")

if __name__ == "__main__":
  main()
